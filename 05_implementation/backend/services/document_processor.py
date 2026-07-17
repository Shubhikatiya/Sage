import os
import re
from pathlib import Path

def extract_text_from_pdf(file_path):
    """Extract text from PDF with tables preserved as markdown."""
    text = ''
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                # Extract text
                page_text = page.extract_text() or ''
                text += f"\n\n--- Page {i+1} ---\n\n"
                text += page_text
                
                # Extract tables
                tables = page.extract_tables()
                for t_idx, table in enumerate(tables):
                    if table:
                        text += f"\n\n[Table {t_idx+1}]\n"
                        # Convert table to markdown
                        for row_idx, row in enumerate(table):
                            clean_row = [str(cell).replace('\n', ' ') if cell else '' for cell in row]
                            text += '| ' + ' | '.join(clean_row) + ' |\n'
                            if row_idx == 0:
                                text += '|' + '|'.join(['---' for _ in row]) + '|\n'
    except Exception as e:
        print(f'pdfplumber error: {e}, falling back to PyPDF2')
        try:
            import PyPDF2
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() or ''
        except Exception as e2:
            print(f'PyPDF2 fallback error: {e2}')
    
    return text.strip()

def extract_text_from_docx(file_path):
    """Extract text from DOCX with tables."""
    text = ''
    try:
        import docx
        doc = docx.Document(file_path)
        
        for element in doc.element.body:
            if element.tag.endswith('tbl'):  # Table
                table = docx.table.Table(element, doc)
                for row in table.rows:
                    cells = [cell.text.replace('\n', ' ') for cell in row.cells]
                    text += '| ' + ' | '.join(cells) + ' |\n'
                text += '\n'
            elif element.tag.endswith('p'):  # Paragraph
                para = docx.text.paragraph.Paragraph(element, doc)
                if para.text.strip():
                    text += para.text + '\n\n'
    except Exception as e:
        print(f'DOCX extraction error: {e}')
        try:
            import docx
            doc = docx.Document(file_path)
            return '\n\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
        except:
            return ''
    
    return text.strip()

def extract_text_from_txt(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            return file.read()
    except UnicodeDecodeError:
        with open(file_path, 'r', encoding='latin-1') as file:
            return file.read()
    except Exception as e:
        print(f'TXT extraction error: {e}')
        return ''

def process_document(file_path):
    """Extract text from various document formats."""
    ext = Path(file_path).suffix.lower()
    
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext == '.docx':
        return extract_text_from_docx(file_path)
    elif ext in ['.txt', '.md']:
        return extract_text_from_txt(file_path)
    else:
        return None

def summarize_document(text, max_length=4000):
    """Create a smart summary of document text."""
    if not text:
        return "No text could be extracted from this document."
    
    lines = text.split('\n')
    paragraphs = []
    current_para = []
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if current_para:
                paragraphs.append(' '.join(current_para))
                current_para = []
        elif stripped.startswith('|') and stripped.endswith('|'):
            # Table row - keep as-is
            if current_para:
                paragraphs.append(' '.join(current_para))
                current_para = []
            paragraphs.append(stripped)
        else:
            current_para.append(stripped)
    
    if current_para:
        paragraphs.append(' '.join(current_para))
    
    # Identify key sections
    summary_parts = []
    total_length = 0
    
    # Add document stats
    word_count = len(text.split())
    summary_parts.append(f"Document Summary ({word_count} words)")
    summary_parts.append("=" * 40)
    
    # Extract headings (lines that look like titles)
    headings = []
    for para in paragraphs[:50]:  # Check first 50 paragraphs
        if len(para) < 100 and para and not para.startswith('|'):
            if para.isupper() or para.endswith(':') or len(para.split()) < 8:
                headings.append(para)
    
    if headings:
        summary_parts.append("\nKey Topics:")
        for h in headings[:10]:
            summary_parts.append(f"- {h}")
    
    # Add first meaningful paragraph as overview
    for para in paragraphs:
        if len(para) > 50 and not para.startswith('---'):
            summary_parts.append(f"\nOverview:\n{para[:500]}")
            break
    
    # Add tables count
    tables = [p for p in paragraphs if p.startswith('|')]
    if tables:
        summary_parts.append(f"\nTables Found: {len(tables)} rows of tabular data preserved.")
    
    summary = '\n'.join(summary_parts)
    
    # Truncate if too long
    if len(summary) > max_length:
        summary = summary[:max_length] + "\n\n[...document continues...]"
    
    return summary
